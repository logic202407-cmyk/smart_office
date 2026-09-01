from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def set_run_font(run, ascii_font: str, east_asia_font: str, size: int | None = None, bold: bool | None = None):
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia_font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_paragraph_spacing(paragraph, before=0, after=0, line=1.4):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines):
        line = lines[i].rstrip()
        if "|" not in line or not line.strip():
            break
        rows.append([c.strip() for c in line.strip().strip("|").split("|")])
        i += 1

    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in rows[1]):
        rows.pop(1)

    return rows, i


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    col_count = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for r_idx, row in enumerate(rows):
        for c_idx in range(col_count):
            text = row[c_idx] if c_idx < len(row) else ""
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(text) <= 12 else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_spacing(p, before=0, after=0, line=1.2)
            run = p.add_run(text)
            set_run_font(run, "Calibri", "宋体", size=10, bold=r_idx == 0)
            if r_idx == 0:
                shade_cell(cell, "D9EAF7")

    doc.add_paragraph()


def add_code_block(doc: Document, code_lines: list[str]):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, "F5F5F5")
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, before=3, after=3, line=1.1)
    for idx, line in enumerate(code_lines):
        run = p.add_run(line)
        set_run_font(run, "Consolas", "等线", size=9)
        if idx != len(code_lines) - 1:
            run.add_break()
    doc.add_paragraph()


def add_paragraph_with_inline(doc: Document, text: str, style: str | None = None, first_line_indent=False):
    p = doc.add_paragraph(style=style)
    if first_line_indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    set_paragraph_spacing(p, before=0, after=6, line=1.45)

    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            set_run_font(run, "Calibri", "黑体", size=11, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            set_run_font(run, "Consolas", "等线", size=10)
        else:
            run = p.add_run(part)
            set_run_font(run, "Calibri", "宋体", size=11)
    return p


def find_first_h1(lines: list[str], fallback_title: str) -> tuple[str, int]:
    for idx, line in enumerate(lines):
        match = re.match(r"^#\s+(.*)$", line.strip())
        if match:
            return match.group(1).strip(), idx
    return fallback_title, -1


def build_doc(markdown_path: Path, output_path: Path):
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    doc_title, first_h1_index = find_first_h1(lines, markdown_path.stem)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(doc_title)
    set_run_font(title_run, "Calibri", "黑体", size=18, bold=True)
    title_run.font.color.rgb = RGBColor(31, 78, 121)
    set_paragraph_spacing(title, before=0, after=12, line=1.0)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run("Markdown 转 Word 预览版")
    set_run_font(sub_run, "Calibri", "等线", size=10)
    sub_run.italic = True
    set_paragraph_spacing(subtitle, before=0, after=18, line=1.0)

    i = 0
    in_code = False
    code_lines: list[str] = []
    while i < len(lines):
        line = lines[i].rstrip("\n")
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            doc.add_paragraph()
            i += 1
            continue

        if "|" in stripped and stripped.count("|") >= 2:
            rows, next_i = parse_table(lines, i)
            if rows:
                add_table(doc, rows)
                i = next_i
                continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if level == 1 and i == first_h1_index and text == doc_title:
                i += 1
                continue
            p = doc.add_paragraph()
            p.style = f"Heading {min(level, 4)}"
            set_paragraph_spacing(p, before=10 if level <= 2 else 6, after=6, line=1.2)
            run = p.add_run(text)
            size = {1: 16, 2: 14, 3: 12, 4: 11}.get(level, 11)
            set_run_font(run, "Calibri", "黑体", size=size, bold=True)
            if level <= 2:
                run.font.color.rgb = RGBColor(31, 78, 121)
            i += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            text = re.sub(r"^\d+\.\s+", "", stripped)
            p = add_paragraph_with_inline(doc, text)
            p.style = "List Number"
            p.paragraph_format.left_indent = Cm(0.63)
            p.paragraph_format.first_line_indent = Cm(0)
            i += 1
            continue

        if stripped.startswith("- "):
            text = stripped[2:].strip()
            p = add_paragraph_with_inline(doc, text)
            p.style = "List Bullet"
            p.paragraph_format.left_indent = Cm(0.63)
            p.paragraph_format.first_line_indent = Cm(0)
            i += 1
            continue

        add_paragraph_with_inline(doc, stripped, first_line_indent=True)
        i += 1

    if in_code and code_lines:
        add_code_block(doc, code_lines)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main():
    if len(sys.argv) != 3:
        print("usage: md_to_docx_preview.py <input.md> <output.docx>")
        sys.exit(1)
    build_doc(Path(sys.argv[1]), Path(sys.argv[2]))


if __name__ == "__main__":
    main()
